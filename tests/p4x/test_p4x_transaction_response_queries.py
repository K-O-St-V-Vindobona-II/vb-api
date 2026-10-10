"""Transaction lists and the summary export look partners and filter hits up
in bulk: the number of statements does not grow with the number of rows."""

import uuid
from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import event, text

from app.api.router_includes import p4x as p4x_routes
from app.models.contact import Contact
from app.models.enums import SubjectMode
from app.models.member import Member
from app.models.org import Org
from app.models.p4x_category import P4xCategory
from app.models.p4x_category_filter import P4xCategoryFilter
from app.models.p4x_partner import P4xPartner
from app.models.state import State
from app.services import (
    p4x_account_service,
    p4x_category_service,
    p4x_partner_service,
    p4x_response_builders,
)
from app.services.p4x_response_builders import (
    build_transaction_response,
    build_transaction_responses,
)
from app.services.p4x_summary_service import generate_summary_xlsx
from tests.p4x.test_p4x_service_queries import _add_tx, _create_account, _now


def _setup(db_session):
    account = _create_account(db_session)
    db_session.add_all(
        [
            Org(id="vbw", label="V", sort_order=1),
            State(id="up", label="U", sort_order=1),
        ]
    )
    db_session.commit()
    category = P4xCategory(
        name="c1",
        label="c1",
        background_color="#000000",
        text_color="#ffffff",
        created_at=_now(),
        updated_at=_now(),
    )
    db_session.add(category)
    db_session.commit()
    category_filter = P4xCategoryFilter(
        name="f",
        p4x_account_id=account.id,
        subject="Test",
        subject_mode=SubjectMode.CONTAINS,
        p4x_category_id=category.id,
        created_at=_now(),
        updated_at=_now(),
    )
    db_session.add(category_filter)
    db_session.commit()
    return account, category_filter


def _add_rows(db_session, account, category_filter, n, tag):
    """Adds `n` bookings, each with a partner member and a hit of the filter."""
    for i in range(n):
        member = Member(vorname=f"M{tag}{i}", nachname="X", org_id="vbw", state_id="up")
        db_session.add(member)
        db_session.commit()
        db_session.add(
            P4xPartner(
                iban=f"AT{tag}{i}",
                member_id=member.id,
                created_at=_now(),
                updated_at=_now(),
            )
        )
        db_session.commit()
        _add_tx(
            db_session,
            account,
            date(2026, 3, 10),
            10.0,
            iban=f"AT{tag}{i}",
            hash_suffix=f"{tag}{i}",
        )
    p4x_category_service.apply_single_filter(db_session, category_filter)


def _month_items(db_session, account):
    items, _ = p4x_account_service.get_transactions_by_month(
        db_session, account, 2026, 3, 1
    )
    return items


def _count(db_session, count_queries, account):
    # The test session starts its savepoint with its first statement; run that
    # one before counting so it is not counted for one measurement only.
    db_session.execute(text("SELECT 1"))
    db_session.expire_all()
    with count_queries() as counter:
        responses = build_transaction_responses(
            _month_items(db_session, account), db_session
        )
    return len(responses), counter.count


class TestTransactionResponsesInBulk:
    def test_statement_count_does_not_grow_with_rows(self, db_session, count_queries):
        account, category_filter = _setup(db_session)
        _add_rows(db_session, account, category_filter, 2, "a")
        small_rows, small = _count(db_session, count_queries, account)
        _add_rows(db_session, account, category_filter, 8, "b")

        large_rows, large = _count(db_session, count_queries, account)

        assert (small_rows, large_rows) == (2, 10)
        assert large == small

    def test_statement_count_does_not_grow_with_delegating_partners(
        self, db_session, count_queries
    ):
        account, category_filter = _setup(db_session)
        _add_rows(db_session, account, category_filter, 2, "a")
        small_rows, small = _count(db_session, count_queries, account)
        for i in range(8):
            member = Member(vorname=f"D{i}", nachname="X", org_id="vbw", state_id="up")
            db_session.add(member)
            db_session.commit()
            _add_tx(
                db_session,
                account,
                date(2026, 3, 11),
                5.0,
                iban=f"ATD{i}",
                hash_suffix=f"d{i}",
                delegating_member_id=member.id,
            )

        large_rows, large = _count(db_session, count_queries, account)

        assert (small_rows, large_rows) == (2, 10)
        assert large == small

    def test_every_row_carries_its_own_partner_and_hit_count(self, db_session):
        account, category_filter = _setup(db_session)
        _add_rows(db_session, account, category_filter, 3, "c")

        responses = build_transaction_responses(
            _month_items(db_session, account), db_session
        )

        assert sorted(r.partner.cn for r in responses) == [
            "Mc0 X",
            "Mc1 X",
            "Mc2 X",
        ]
        assert {r.partner.type for r in responses} == {"member"}
        assert [r.p4x_category_filters[0].hitCount for r in responses] == [3, 3, 3]

    def test_the_single_row_form_gives_the_same_response(self, db_session):
        account, category_filter = _setup(db_session)
        _add_rows(db_session, account, category_filter, 3, "c")
        items = _month_items(db_session, account)

        single = build_transaction_response(items[0], db_session)
        bulk = build_transaction_responses(items, db_session)[0]

        assert single == bulk
        assert single.partner is not None
        assert single.partner.type == "member"
        assert single.p4x_category_filters[0].hitCount == 3

    def test_the_delegating_partner_is_resolved(self, db_session):
        account, _ = _setup(db_session)
        contact = Contact(
            name="Delegiert",
            kontakttyp="person",
            org_id="vbw",
            zustellungen=True,
        )
        db_session.add(contact)
        db_session.commit()
        tx = _add_tx(db_session, account, date(2026, 3, 10), 10.0)
        tx.delegating_contact_id = contact.id
        db_session.commit()

        response = build_transaction_response(tx, db_session)

        assert response.delegating_partner is not None
        assert response.delegating_partner.type == "contact"
        assert response.delegating_partner.cn == "Delegiert"

    def test_a_soft_deleted_partner_gives_no_partner(self, db_session):
        account, _ = _setup(db_session)
        member = Member(vorname="Weg", nachname="X", org_id="vbw", state_id="up")
        db_session.add(member)
        db_session.commit()
        partner = P4xPartner(
            iban="ATGONE",
            member_id=member.id,
            created_at=_now(),
            updated_at=_now(),
        )
        db_session.add(partner)
        db_session.commit()
        tx = _add_tx(db_session, account, date(2026, 3, 10), 10.0, iban="ATGONE")
        assert build_transaction_response(tx, db_session).partner is not None

        partner.deleted_at = _now()
        db_session.commit()
        db_session.expire_all()

        assert build_transaction_response(tx, db_session).partner is None

    def test_a_row_without_a_filter_hit_has_no_filters(self, db_session):
        account, _ = _setup(db_session)
        tx = _add_tx(db_session, account, date(2026, 3, 10), 10.0, subject="Other")

        response = build_transaction_response(tx, db_session)

        assert response.p4x_category_filters == []


class TestFindPartnerEntities:
    def test_unknown_types_are_ignored(self, db_session):
        _setup(db_session)
        member = Member(vorname="A", nachname="B", org_id="vbw", state_id="up")
        db_session.add(member)
        db_session.commit()

        found = p4x_partner_service.find_partner_entities(
            db_session, [("member", member.id), ("nonsense", member.id)]
        )

        assert list(found) == [("member", member.id)]

    def test_missing_rows_are_absent(self, db_session):
        _setup(db_session)

        found = p4x_partner_service.find_partner_entities(
            db_session, [("member", uuid.uuid4())]
        )

        assert found == {}

    def test_no_references_cost_no_statement(self, db_session, count_queries):
        with count_queries() as counter:
            found = p4x_partner_service.find_partner_entities(db_session, [])

        assert found == {}
        assert counter.count == 0

    def test_one_statement_per_partner_type(self, db_session, count_queries):
        _setup(db_session)
        members = [
            Member(vorname=f"A{i}", nachname="B", org_id="vbw", state_id="up")
            for i in range(5)
        ]
        contact = Contact(
            name="K", kontakttyp="person", org_id="vbw", zustellungen=True
        )
        db_session.add_all([*members, contact])
        db_session.commit()
        refs = [("member", m.id) for m in members] + [("contact", contact.id)]
        db_session.expire_all()

        with count_queries() as counter:
            found = p4x_partner_service.find_partner_entities(db_session, refs)

        assert len(found) == 6
        assert counter.count == 2


class TestSummaryExportPartnerLookups:
    def test_member_reads_do_not_grow_with_the_number_of_bookings(self, db_session):
        account, category_filter = _setup(db_session)
        _add_rows(db_session, account, category_filter, 2, "s")
        small = self._member_reads(db_session)
        _add_rows(db_session, account, category_filter, 8, "t")

        large = self._member_reads(db_session)

        assert small >= 1
        assert large == small

    @staticmethod
    def _member_reads(db_session):
        """Number of statements that read the members table during one export."""
        recorded: list[str] = []

        def on_execute(_conn, _cursor, statement, *_args):
            if "FROM members" in statement:
                recorded.append(statement)

        db_session.execute(text("SELECT 1"))
        db_session.expire_all()
        bind = db_session.get_bind()
        event.listen(bind, "before_cursor_execute", on_execute)
        try:
            generate_summary_xlsx(db_session, date(2026, 3, 1), date(2026, 3, 31))
        finally:
            event.remove(bind, "before_cursor_execute", on_execute)
        return len(recorded)

    def test_the_partner_column_names_member_contact_and_unknown(self, db_session):
        account, _ = _setup(db_session)
        member = Member(vorname="Max", nachname="Muster", org_id="vbw", state_id="up")
        contact = Contact(
            name="Delegiert", kontakttyp="person", org_id="vbw", zustellungen=True
        )
        db_session.add_all([member, contact])
        db_session.commit()
        db_session.add(
            P4xPartner(
                iban="ATM",
                member_id=member.id,
                created_at=_now(),
                updated_at=_now(),
            )
        )
        db_session.commit()
        _add_tx(db_session, account, date(2026, 3, 1), 1.0, iban="ATM", hash_suffix="1")
        _add_tx(db_session, account, date(2026, 3, 2), 2.0, iban="ATX", hash_suffix="2")
        delegated = _add_tx(
            db_session, account, date(2026, 3, 3), 3.0, iban="ATM", hash_suffix="3"
        )
        delegated.delegating_contact_id = contact.id
        db_session.commit()

        raw, _ = generate_summary_xlsx(db_session, date(2026, 3, 1), date(2026, 3, 31))

        sheet = load_workbook(BytesIO(raw))[account.label]
        partners = [sheet.cell(row=r, column=2).value for r in (2, 3, 4)]
        assert partners == [
            "Mitglied: Max Muster",
            "unknown: unknown",
            "Kontakt: Delegiert",
        ]


class TestListEndpointsUseTheBulkBuilder:
    """Every list endpoint hands its whole page to the bulk builder; calling
    the single-row builder per item would bring the extra statements back."""

    @pytest.fixture
    def builder_calls(self, monkeypatch):
        calls: dict[str, list[int]] = {"bulk": [], "single": []}
        real_bulk = p4x_response_builders.build_transaction_responses

        def bulk(txs, db):
            calls["bulk"].append(len(txs))
            return real_bulk(txs, db)

        def single(tx, db):
            calls["single"].append(1)
            return real_bulk([tx], db)[0]

        monkeypatch.setattr(p4x_response_builders, "build_transaction_responses", bulk)
        monkeypatch.setattr(p4x_response_builders, "build_transaction_response", single)
        return calls

    @staticmethod
    def _data(db_session):
        account, category_filter = _setup(db_session)
        _add_rows(db_session, account, category_filter, 3, "r")
        # A booking without a partner and without a category shows up in
        # both warning lists.
        _add_tx(db_session, account, date(2026, 3, 12), 7.0, iban="ATNONE", subject="Z")
        member = db_session.query(Member).filter_by(vorname="Mr0").one()
        return account, category_filter, member

    @pytest.mark.parametrize(
        "endpoint",
        [
            "dashboard",
            "warnings_partner",
            "warnings_category",
            "by_month",
            "by_partner",
            "by_category",
            "by_filter",
        ],
    )
    def test_the_endpoint_builds_its_page_in_bulk(
        self, db_session, builder_calls, endpoint
    ):
        account, category_filter, member = self._data(db_session)
        calls = {
            "dashboard": lambda: p4x_routes.get_dashboard(db_session, member),
            "warnings_partner": lambda: p4x_routes.get_warnings_partner_list(
                db_session, member
            ),
            "warnings_category": lambda: p4x_routes.get_warnings_category_list(
                db_session, member
            ),
            "by_month": lambda: p4x_routes.get_transactions_by_month(
                account.id, 2026, 3, db_session, member
            ),
            "by_partner": lambda: p4x_routes.get_transactions_by_partner(
                account.id, "member", member.id, db_session, member
            ),
            "by_category": lambda: p4x_routes.get_transactions_by_category(
                account.id, category_filter.p4x_category_id, db_session, member
            ),
            "by_filter": lambda: p4x_routes.get_transactions_by_filter(
                account.id, category_filter.id, db_session, member
            ),
        }

        calls[endpoint]()

        pages = 2 if endpoint == "dashboard" else 1
        assert builder_calls["single"] == []
        assert len(builder_calls["bulk"]) == pages
        assert min(builder_calls["bulk"]) >= 1
